"""DeckWright tests: runtime, presenter view, ink, print, player and CLI in real Chromium (file://).
Run: npm test   (or: node tools/build.mjs && python3 test/run_tests.py)
Needs: pip install playwright pdfplumber python-pptx; npm i -g playwright pptxgenjs (for export)
"""
import os, sys, subprocess, tempfile, json, zipfile
from playwright.sync_api import sync_playwright

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIST = os.path.join(ROOT, 'dist')
TMP = tempfile.mkdtemp(prefix='deckwright-test-')
results = []

def check(name, cond, detail=''):
    results.append((name, bool(cond), detail))
    print(('  ✓ ' if cond else '  ✗ ') + name + ('' if cond else f'  → {detail}'))

def cli(*args):
    return subprocess.run(['node', os.path.join(ROOT, 'skills', 'deckwright', 'scripts', 'deckwright.mjs'), *args], cwd=ROOT, capture_output=True, text=True)

def url(p): return 'file://' + os.path.abspath(p)

with sync_playwright() as pw:
    browser = pw.chromium.launch()

    # ------------------------------------------------------------ single HTML
    print('Single HTML (dist/demo.html)')
    ctx = browser.new_context(viewport={'width': 1280, 'height': 720})
    pg = ctx.new_page(); errors = []
    pg.on('pageerror', lambda e: errors.append(str(e)))
    pg.goto(url(os.path.join(DIST, 'demo.html'))); pg.wait_for_timeout(500)
    st = lambda: pg.evaluate('DeckWright.state()')
    s = st()
    check('contract mode, 7 slides', s['mode'] == 'contract' and s['count'] == 7, s)
    check('1920×1080 stage scaled to fit', 'scale(' in pg.evaluate("document.querySelector('[data-hdek]').style.transform"))
    for _ in range(3): pg.keyboard.press('ArrowRight'); pg.wait_for_timeout(120)
    check('→ ×3: slide 2, step 2', st()['slide'] == 1 and st()['step'] == 2, st())
    pg.keyboard.press('ArrowLeft'); pg.wait_for_timeout(30)
    ops = pg.evaluate("[...document.querySelectorAll('.hdek-current .hdek-step')].map(e=>getComputedStyle(e).opacity)")
    check('← snaps to previous step without animation', ops[:2] == ['1', '0'], ops)
    pg.keyboard.type('3'); pg.keyboard.press('Enter'); pg.wait_for_timeout(100)
    pg.keyboard.press('ArrowRight'); pg.wait_for_timeout(1500)
    check('number+Enter jumps to slide 3, JS hook counter reaches 128', pg.evaluate("document.querySelector('[data-count]').textContent") == '128', pg.evaluate("document.querySelector('[data-count]').textContent"))
    pg.keyboard.press('ArrowRight'); pg.wait_for_timeout(100); pg.keyboard.press('ArrowLeft'); pg.wait_for_timeout(50)
    check('hook slide: back restores state instantly', pg.evaluate("document.querySelector('[data-bar]').style.height") == '0%')
    pg.keyboard.type('5'); pg.keyboard.press('Enter'); pg.wait_for_timeout(150)
    pg.click('.card >> nth=0'); pg.wait_for_timeout(150)
    check('clicking a cursor:pointer element does not advance', st()['slide'] == 4)
    pg.mouse.click(1200, 80); pg.wait_for_timeout(150)
    check('clicking empty space advances', st()['slide'] == 5)
    pg.mouse.wheel(0, 150); pg.wait_for_timeout(150)
    check('wheel down = next step', st()['step'] == 1, st())
    pg.keyboard.press('b'); check('B = black screen', pg.evaluate("document.querySelector('.hdek-blank').classList.contains('hdek-blank-on')"))
    pg.keyboard.press('b')
    check('position stored in URL hash', pg.evaluate('location.hash') == '#6.1', pg.evaluate('location.hash'))

    # presenter view
    with ctx.expect_page() as pi: pg.keyboard.press('s')
    pp = pi.value; pp.wait_for_timeout(1800)
    frames = {f.name: f for f in pp.frames if f.name.startswith('hdek-preview')}
    check('presenter view opens with 2 previews', len(frames) == 2, list(frames))
    check('presenter slide number', pp.inner_text('#num') == '6 / 7', pp.inner_text('#num'))
    check('current preview matches audience', frames['hdek-preview-cur'].evaluate('DeckWright.state().step') == 1)
    check('next preview shows next step', frames['hdek-preview-nxt'].evaluate('[DeckWright.state().slide, DeckWright.state().step]') == [5, 2])
    check('notes shown', 'data-step-out' in pp.inner_text('#notes'))
    pp.keyboard.press('ArrowRight'); pp.wait_for_timeout(300)
    check('advancing in presenter moves audience', st()['step'] == 2, st())
    pp.click('button[data-act="grid"]'); pp.wait_for_selector('.hdg-card'); pp.click('.hdg-card >> nth=1'); pp.wait_for_timeout(300)
    check('pick slide from the overview (all steps revealed)', st()['slide'] == 1 and st()['step'] == st()['max'], st())
    b0 = st(); pp.click('#curbox'); pp.wait_for_timeout(300)
    check('click on the presenter preview advances', st()['step'] == b0['step'] + 1 or st()['slide'] == b0['slide'] + 1, (b0, st()))
    b1 = st(); pp.click('#notes'); pp.wait_for_timeout(300)
    check('click on the notes does not advance', st() == b1, (b1, st()))
    pp.mouse.move(*[v + 20 for v in (pp.locator('#curbox').bounding_box()['x'], pp.locator('#curbox').bounding_box()['y'])]); pp.mouse.wheel(0, -120); pp.wait_for_timeout(300)
    check('wheel up over the preview goes back', st() == b0, (b0, st()))
    # ink drawn in the presenter view is stored on the audience side
    pp.click('button[data-act="pen"]')
    box = pp.locator('#curbox').bounding_box()
    pp.mouse.move(box['x'] + 50, box['y'] + 60); pp.mouse.down()
    for k in range(10): pp.mouse.move(box['x'] + 50 + k * 20, box['y'] + 60 + k * 6)
    pp.mouse.up(); pp.wait_for_timeout(200)
    ink = pg.evaluate('DeckWright.ink()')
    check('presenter ink saved in audience data', len(ink.get('steps', [])) == 1, ink.keys())
    px = pg.evaluate("(()=>{const c=document.querySelector('.hdek-ink');const d=c.getContext('2d').getImageData(0,0,c.width,c.height).data;let n=0;for(let i=3;i<d.length;i+=4)if(d[i])n++;return n})()")
    check('ink rendered on audience canvas', px > 100, px)
    pp.keyboard.press('Escape')
    pg.reload(); pg.wait_for_timeout(500)
    check('position and ink survive reload', st()['slide'] == 1 and len(pg.evaluate('DeckWright.ink()').get('steps', [])) == 1, st())
    # print
    pdf = os.path.join(TMP, 'print.pdf')
    pg.pdf(path=pdf, prefer_css_page_size=True, print_background=True)
    import pdfplumber
    with pdfplumber.open(pdf) as d: n = len(d.pages)
    check('Ctrl+P print = one page per slide', n == 7, n)
    check('state restored after print', st()['slide'] == 1 and pg.evaluate("getComputedStyle(document.querySelectorAll('.hdek-slide')[0]).display") == 'none')
    check('no page errors', not errors, errors)
    ctx.close()

    # UI language follows the browser (ko / en)
    for loc, word in (('ko-KR', '다음'), ('en-US', 'Next')):
        ctx = browser.new_context(viewport={'width': 1280, 'height': 720}, locale=loc); pg = ctx.new_page()
        pg.goto(url(os.path.join(DIST, 'demo.html'))); pg.wait_for_timeout(400)
        pg.keyboard.press('?'); pg.wait_for_timeout(100)
        check(f'UI language follows browser ({loc})', word in pg.inner_text('.hdek-help'), pg.inner_text('.hdek-help')[:80])
        ctx.close()

    # ------------------------------------------------------------ deck without contract
    print('AI-style deck without contract (heuristic)')
    ctx = browser.new_context(viewport={'width': 1280, 'height': 720}); pg = ctx.new_page()
    pg.goto(url(os.path.join(DIST, 'ai-style.html'))); pg.wait_for_timeout(500)
    s = pg.evaluate('DeckWright.state()')
    check('heuristic mode, 4 slides', s['mode'] == 'heuristic' and s['count'] == 4, s)
    check('.notes used as speaker notes', 'knows nothing about DeckWright' in pg.evaluate('DeckWright.slides()[0].notes'))
    for _ in range(3): pg.keyboard.press('ArrowRight'); pg.wait_for_timeout(120)
    s = pg.evaluate('DeckWright.state()')
    check('.fragment as steps (slide 3, step 1)', s['slide'] == 2 and s['step'] == 1, s)
    check('no double move with the deck\'s own key handler', pg.evaluate("document.querySelectorAll('.slide.active').length") == 1)
    ctx.close()

    # ------------------------------------------------------------ deck with its own engine (driver)
    print('Deck with its own navigation engine (driver)')
    out = os.path.join(TMP, 'stage.html'); cli('build', 'test/fixtures/stage-engine', '-o', out)
    ctx = browser.new_context(viewport={'width': 1280, 'height': 720}); pg = ctx.new_page(); errors = []
    pg.on('pageerror', lambda e: errors.append(str(e)))
    pg.goto(url(out)); pg.wait_for_function('window.DeckWright && DeckWright.state().count>1')
    st = lambda: pg.evaluate('[DeckWright.state().slide, DeckWright.state().step]')
    check('driver mode, 4 scenes', pg.evaluate('DeckWright.state().mode') == 'driver' and pg.evaluate('DeckWright.state().count') == 4)
    check('aside.note inside scene used as notes', pg.evaluate('DeckWright.slides()[3].notes') == '끝 노트')
    seq = []
    for _ in range(4): pg.keyboard.press('ArrowRight'); pg.wait_for_timeout(80); seq.append(st())
    check('→ handled by deck engine, state tracked', seq == [[0, 1], [0, 2], [1, 0], [1, 1]], seq)
    check('matches deck\'s own counter', pg.inner_text('#count') == '2/4')
    pg.keyboard.press('ArrowDown'); pg.wait_for_timeout(500)  # the deck gets ~300 ms to react first
    check('↓ (unknown to deck) handled by runtime', st() == [2, 0], st())
    pg.mouse.click(900, 400); pg.wait_for_timeout(500)
    check('click handled by deck: no double move', st() == [3, 0], st())
    pg.keyboard.type('4'); pg.keyboard.press('Enter'); pg.wait_for_timeout(150)
    check('number jump: deck does not also handle Enter', st() == [3, 0], st())
    pg.evaluate('DeckWright.goto(3, 2, false)'); pg.wait_for_timeout(100)
    check('goto(slide, step) snaps', st() == [3, 2], st())
    with ctx.expect_page() as pi: pg.keyboard.press('s')
    pp = pi.value; pp.wait_for_timeout(1800)
    fr = {f.name: f for f in pp.frames if f.name.startswith('hdek-preview')}
    check('driver presenter previews in sync', fr['hdek-preview-cur'].evaluate('[DeckWright.state().slide,DeckWright.state().step]') == [3, 2] and fr['hdek-preview-nxt'].evaluate('[DeckWright.state().slide,DeckWright.state().step]') == [3, 3])
    check('previews do not change audience state', st() == [3, 2], st())
    pdf = os.path.join(TMP, 'stage.pdf'); pg.pdf(path=pdf, prefer_css_page_size=True, print_background=True)
    import pdfplumber
    with pdfplumber.open(pdf) as d: n = len(d.pages)
    check('driver deck print = one page per scene', n == 4, n)
    check('driver: no page errors', not errors, errors)
    ctx.close()

    # ------------------------------------------------------------ stacked fixed-size slides (autostage)
    print('Stacked fixed-size slides (autostage)')
    out = os.path.join(TMP, 'stacked.html'); cli('build', 'test/fixtures/stacked-pt', '-o', out)
    ctx = browser.new_context(viewport={'width': 1280, 'height': 720}); pg = ctx.new_page()
    pg.goto(url(out)); pg.wait_for_function('window.DeckWright && DeckWright.state().count>1')
    check('starts after deck\'s load handlers (measurements non-zero)', pg.evaluate("[...document.querySelectorAll('.fit')].every(e=>e.textContent==='measured')"))
    check('sibling aside.notes used as notes', pg.evaluate('DeckWright.slides()[1].notes').find('노트 2') >= 0)
    r = pg.evaluate("(()=>{const r=document.querySelector('.hdek-current').getBoundingClientRect();return [r.left,r.top,r.width,r.height]})()")
    check('720pt×405pt slide fills the viewport', abs(r[2] - 1280) < 2 and abs(r[3] - 720) < 2, r)
    check('notes hidden on screen', pg.evaluate("getComputedStyle(document.querySelector('aside.notes')).display") == 'none')
    ctx.close()

    # ------------------------------------------------------------ player
    print('player.html + .hdek')
    ctx = browser.new_context(viewport={'width': 1280, 'height': 720}); pg = ctx.new_page(); errors = []
    pg.on('pageerror', lambda e: errors.append(str(e)))
    pg.goto(url(os.path.join(DIST, 'player.html')))
    pg.set_input_files('#file', os.path.join(DIST, 'demo.hdek')); pg.wait_for_selector('#info', state='visible')
    check('package info shown', ('7 slides' in pg.inner_text('#meta') or '슬라이드 7개' in pg.inner_text('#meta')), pg.inner_text('#meta'))
    pg.click('#start'); pg.wait_for_timeout(800)
    check('start presenting from player', pg.evaluate('DeckWright.state().count') == 7)
    with ctx.expect_page() as pi: pg.keyboard.press('s')
    pp = pi.value; pp.wait_for_timeout(1800)
    fr = [f for f in pp.frames if f.name.startswith('hdek-preview')]
    check('player-mode presenter previews work', all(f.evaluate('!!window.DeckWright') for f in fr) and len(fr) == 2)
    check('player: no page errors', not errors, errors)
    # open folder
    pg2 = ctx.new_page(); pg2.goto(url(os.path.join(DIST, 'player.html')))
    fx = os.path.join(ROOT, 'test', 'fixtures', 'assets-deck')
    pg2.set_input_files('#dir', fx); pg2.wait_for_selector('#info', state='visible')
    check('open folder: warns about external web font', 'fonts.googleapis.com' in pg2.inner_text('#issues'))
    pg2.click('#start'); pg2.wait_for_timeout(800)
    pg3 = ctx.new_page(); pg3.goto(url(os.path.join(DIST, 'player.html')))
    pg3.set_input_files('#file', os.path.join(DIST, 'demo.hdek')); pg3.wait_for_selector('#info', state='visible')
    bad = os.path.join(TMP, 'bad.hdek'); open(bad, 'wb').write(open(os.path.join(DIST, 'demo.hdek'), 'rb').read()[:4000])
    pg3.set_input_files('#file', bad); pg3.wait_for_timeout(600)
    check('player: a failed load hides the previous deck', not pg3.is_visible('#info') and pg3.is_visible('#err'), pg3.inner_text('#err'))
    check('folder deck: CSS fonts/backgrounds inlined', pg2.evaluate("getComputedStyle(document.querySelector('section')).backgroundImage.startsWith('url(\"data:image/png')"))
    ctx.close()

    # ------------------------------------------------------------ regressions from the code review
    print('Review regressions (runtime)')
    def built(fx):
        out = os.path.join(TMP, fx + '.html'); r = cli('build', 'test/fixtures/' + fx, '-o', out)
        assert r.returncode == 0, r.stderr
        return out
    # driver deck that ignores keys during its 500 ms transition: number+Enter must still land on slide 4
    ctx = browser.new_context(viewport={'width': 1280, 'height': 720}); pg = ctx.new_page()
    pg.goto(url(built('lock-engine'))); pg.wait_for_function('window.DeckWright && DeckWright.state().count>1')
    pg.keyboard.type('4'); pg.keyboard.press('Enter'); pg.wait_for_timeout(6000)
    st = pg.evaluate('[DeckWright.state().slide, DeckWright.state().step]')
    check('driver jump waits for a deck that locks keys during transitions', st[0] == 3, st)
    ctx.close()
    # driver deck that navigates on keyup: one press = one step
    ctx = browser.new_context(viewport={'width': 1280, 'height': 720}); pg = ctx.new_page()
    pg.goto(url(built('keyup-engine'))); pg.wait_for_function('window.DeckWright && DeckWright.state().count>1')
    seq = []
    for _ in range(3):
        pg.keyboard.down('ArrowRight'); pg.wait_for_timeout(90); pg.keyboard.up('ArrowRight'); pg.wait_for_timeout(450)
        seq.append(pg.evaluate('[DeckWright.state().slide, DeckWright.state().step]'))
    check('keyup-driven deck: no double advance', seq == [[0, 1], [0, 2], [1, 0]], seq)
    ctx.close()
    # heuristic deck with its own document click handler: one click = one move, slide stays visible
    ctx = browser.new_context(viewport={'width': 1280, 'height': 720}); pg = ctx.new_page()
    pg.goto(url(built('ai-click'))); pg.wait_for_timeout(800)
    for _ in range(3): pg.mouse.click(1100, 650); pg.wait_for_timeout(700)
    vis = pg.evaluate('(()=>{const c=document.querySelector(".hdek-current");return [DeckWright.state().slide, getComputedStyle(c).opacity, document.querySelectorAll(".slide.active").length]})()')
    check('heuristic click is not handled twice (deck click handler)', vis[1] == '1' and vis[2] == 1, vis)
    # one trackpad fling (momentum) = one step
    s0 = pg.evaluate('JSON.stringify(DeckWright.state())')
    for k in range(60): pg.mouse.wheel(0, max(1, 60 - k)); pg.wait_for_timeout(30)
    s1 = pg.evaluate('DeckWright.state()'); s0 = json.loads(s0)
    check('one trackpad fling advances once', (s1['slide'] * 10 + s1['step']) - (s0['slide'] * 10 + s0['step']) in (1, 10 - s0['step']), (s0, s1))
    ctx.close()
    # a deck that goes back when its left third is clicked: any click moves forward, right-click and wheel-up go back
    ctx = browser.new_context(viewport={'width': 1280, 'height': 720}); pg = ctx.new_page()
    pg.goto(url(built('zone-engine'))); pg.wait_for_function('window.DeckWright && DeckWright.state().count>1')
    zs = lambda: pg.evaluate('[DeckWright.state().slide, DeckWright.state().step]')
    pg.mouse.click(1000, 360); pg.wait_for_timeout(400); a1 = zs()
    pg.mouse.click(150, 360); pg.wait_for_timeout(400); a2 = zs()
    pg.mouse.click(640, 360, button='right'); pg.wait_for_timeout(400); a3 = zs()
    pg.mouse.move(640, 360); pg.mouse.wheel(0, 120); pg.wait_for_timeout(600); a4 = zs()
    pg.mouse.wheel(0, -120); pg.wait_for_timeout(600); a5 = zs()
    check('click anywhere goes forward (even the left side of a deck that would go back)', a1 == [0, 1] and a2 == [0, 2], (a1, a2))
    check('right-click and wheel-up go back, wheel-down forward', a3 == [0, 1] and a4 == [0, 2] and a5 == [0, 1], (a3, a4, a5))
    ctx.close()

    # ------------------------------------------------------------ v0.3.4: resume, overview + search, magnifier, video from the presenter
    print('Resume, overview, magnifier, video control')
    media = built('media-deck')
    ctx = browser.new_context(viewport={'width': 1280, 'height': 720}); pg = ctx.new_page(); errors = []
    pg.on('pageerror', lambda e: errors.append(str(e)))
    ms = lambda: pg.evaluate('[DeckWright.state().slide, DeckWright.state().step]')
    sr = "document.querySelector('.hdek-ui-host').shadowRoot"
    pg.goto(url(media)); pg.wait_for_timeout(400)
    pg.keyboard.type('4'); pg.keyboard.press('Enter'); pg.wait_for_timeout(200)
    pg.goto('about:blank'); pg.goto(url(media)); pg.wait_for_timeout(500)
    toast_txt = pg.evaluate(sr + ".querySelector('.hdek-toast').textContent")
    check('reopening starts at slide 1 and offers the last position', ms() == [0, 0] and '4' in toast_txt, (ms(), toast_txt))
    pg.keyboard.press('r'); pg.wait_for_timeout(300)
    check('R continues where the talk stopped', ms() == [3, 0], ms())
    pg.keyboard.press('Home'); pg.wait_for_timeout(200)
    pg.goto('about:blank'); pg.goto(url(media)); pg.wait_for_timeout(500)
    pg.keyboard.press('r'); pg.wait_for_timeout(200)
    check('no offer (and R does nothing) when the talk ended on slide 1', ms() == [0, 0], ms())
    # overview + search on the audience screen
    pg.keyboard.press('g'); pg.wait_for_timeout(2500)
    nc = pg.evaluate(sr + ".querySelectorAll('.hdg-card').length"); nth = pg.evaluate(sr + ".querySelectorAll('.hdg-th iframe.on').length")
    check('G: overview with a live thumbnail of every slide', nc == 5 and nth == 5, (nc, nth))
    pg.keyboard.type('odi b'); pg.wait_for_timeout(200)
    found = pg.evaluate(sr + ".querySelector('.hdg-n').textContent"); snip = pg.evaluate(sr + ".querySelector('.hdg-card:not([hidden]) .hdg-snip').innerHTML")
    check('search finds slide text; typing letters does not trigger shortcuts', found.startswith('1') and '<mark>ODI</mark>' in snip and not pg.evaluate("document.querySelector('.hdek-blank').classList.contains('hdek-blank-on')"), (found, snip))
    pg.keyboard.press('Enter'); pg.wait_for_timeout(300)
    check('Enter jumps to the match with every step revealed', ms() == [2, 2] and not pg.evaluate(sr + ".querySelector('.hdg')"), ms())
    # magnifier
    pg.keyboard.press('z'); pg.mouse.move(900, 500); pg.mouse.down(); pg.mouse.move(1000, 560); pg.mouse.move(1150, 640); pg.mouse.up(); pg.wait_for_timeout(450)
    zt = pg.evaluate('document.documentElement.style.transform')
    check('Z + drag zooms the projector into that area (no slide change)', 'scale(5' in zt and ms() == [2, 2], (zt, ms()))
    pg.mouse.click(640, 360, button='right'); pg.wait_for_timeout(450)
    check('right-click zooms back out by half', 'scale(2.5' in pg.evaluate('document.documentElement.style.transform'), pg.evaluate('document.documentElement.style.transform'))
    pg.keyboard.press('Escape'); pg.wait_for_timeout(450)
    check('Esc: magnifier off and zoom reset', pg.evaluate('document.documentElement.style.transform') == '' and pg.evaluate('DeckWright.__debug.tool') is None)
    pg.keyboard.press('z'); pg.mouse.click(640, 360); pg.wait_for_timeout(300); pg.keyboard.press('z'); pg.keyboard.press('ArrowLeft'); pg.wait_for_timeout(300)
    check('changing slides clears the zoom', pg.evaluate('document.documentElement.style.transform') == '', pg.evaluate('document.documentElement.style.transform'))
    # video on the projector, controlled from the presenter view
    pg.keyboard.type('2'); pg.keyboard.press('Enter'); pg.wait_for_timeout(200)
    with ctx.expect_page() as pi: pg.keyboard.press('s')
    pp = pi.value; pp.wait_for_timeout(1800)
    vid = "(v=>[v.paused, +v.currentTime.toFixed(1)])(document.querySelector('.hdek-current video'))"
    cur = [f for f in pp.frames if f.name == 'hdek-preview-cur'][0]
    check('presenter shows video controls on a slide with a video', pp.locator('.mrow').count() == 1)
    pp.click('.mrow [data-m="play"]'); pp.wait_for_timeout(1500)
    a_v, p_v = pg.evaluate(vid), cur.evaluate(vid)
    check('▶ in the presenter plays the projector video; the preview follows, muted', a_v[0] is False and p_v[0] is False and abs(a_v[1] - p_v[1]) < 0.8 and cur.evaluate("document.querySelector('.hdek-current video').muted"), (a_v, p_v))
    pp.click('.mrow [data-m="play"]'); pp.click('.mrow [data-m="fwd"]'); pp.wait_for_timeout(700)
    a_v, p_v = pg.evaluate(vid), cur.evaluate(vid)
    check('⏸ and +10 s reach both, and the time is shown', a_v[0] and p_v[0] and a_v[1] >= 10 and abs(a_v[1] - p_v[1]) < 0.2 and pp.inner_text('.mrow .mt').startswith('0:1'), (a_v, p_v, pp.inner_text('.mrow .mt')))
    pg.keyboard.press('k'); pg.wait_for_timeout(300)
    check('K plays / pauses the video on the slide', pg.evaluate(vid)[0] is False); pg.keyboard.press('k')
    pp.keyboard.press('ArrowRight'); pp.wait_for_timeout(400)
    check('video controls go away on a slide without video', pp.locator('.mrow').count() == 0 and pg.evaluate("document.querySelectorAll('video')[0].paused"))
    # magnifier from the presenter view
    pp.keyboard.press('ArrowRight'); pp.keyboard.press('ArrowRight'); pp.wait_for_timeout(300); pp.keyboard.press('z')
    bx = pp.locator('#curbox').bounding_box()
    pp.mouse.move(bx['x'] + bx['width'] * .6, bx['y'] + bx['height'] * .62); pp.mouse.down(); pp.mouse.move(bx['x'] + bx['width'] * .75, bx['y'] + bx['height'] * .75); pp.mouse.move(bx['x'] + bx['width'] * .9, bx['y'] + bx['height'] * .85); pp.mouse.up(); pp.wait_for_timeout(500)
    za, zp = pg.evaluate('document.documentElement.style.transform'), cur.evaluate('document.documentElement.style.transform')
    check('dragging on the presenter preview zooms the projector and the preview alike', 'scale(3.3' in za and 'scale(3.3' in zp and ms() == [2, 2], (za, zp, ms()))
    pp.keyboard.press('Escape'); pp.wait_for_timeout(300)
    # overview in the presenter view
    pp.keyboard.press('g'); pp.wait_for_timeout(1500); pp.keyboard.type('pelvic'); pp.wait_for_timeout(200)
    check('G in the presenter view: overview and search there, not on the projector', pp.inner_text('.hdg-n').startswith('1') and not pg.evaluate(sr + ".querySelector('.hdg')"), pp.inner_text('.hdg-n'))
    pp.keyboard.press('Enter'); pp.wait_for_timeout(300)
    check('picking in the presenter overview moves the projector', ms() == [3, 0], ms())
    check('v0.3.4 features: no page errors', not errors, errors)
    # the audience toolbar's overview button opens on the projector even while the presenter view is open
    pg.evaluate(sr + ".querySelector('.hdek-bar button[data-act=grid]').click()"); pg.wait_for_timeout(300)
    check('toolbar "all slides" opens on the audience screen with the presenter open', pg.evaluate(sr + ".querySelector('.hdg') !== null") and pp.locator('.hdg').count() == 0)
    pg.keyboard.press('Escape')
    ctx.close()

    # v0.3.5 (review): the overview must not leak clicks, typing or the wheel to a deck with its own handlers
    ctx = browser.new_context(viewport={'width': 1280, 'height': 720}); pg = ctx.new_page()
    pg.goto(url(built('ai-click'))); pg.wait_for_timeout(800)
    prog = lambda: pg.evaluate("document.querySelector('.progress').style.width")
    p0 = prog(); pg.keyboard.press('g'); pg.wait_for_timeout(800)
    pg.keyboard.type('a b'); pg.mouse.move(640, 400); pg.mouse.wheel(0, 300); pg.wait_for_timeout(300)
    p1 = prog()
    pg.evaluate("document.querySelector('.hdek-ui-host').shadowRoot.querySelector('.hdg-q').value=''")
    pg.evaluate("document.querySelector('.hdek-ui-host').shadowRoot.querySelectorAll('.hdg-card')[2].click()"); pg.wait_for_timeout(400)
    act = pg.evaluate("[...document.querySelectorAll('.slide')].map((s,i)=>s.classList.contains('active')?i:-1).filter(i=>i>=0)")
    check('overview: typing, wheel and clicks never reach the deck underneath', p0 == p1 and len(act) <= 1 and pg.evaluate('DeckWright.state().slide') == 2, (p0, p1, act, pg.evaluate('DeckWright.state().slide')))
    ctx.close()

    # v0.3.5 (review): zoom on a scrolled page lands on the clicked point, toolbar stays on screen, scroll comes back
    longp = os.path.join(TMP, 'long.html')
    with open(longp, 'w') as f:
        f.write('<!doctype html><title>Long</title><body style="margin:0;height:4000px;font:20px sans-serif"><p>A long page</p>'
                '<div id="mk" style="position:absolute;left:600px;top:2040px;width:40px;height:40px;background:red"></div>'
                '<script src="' + os.path.join(ROOT, 'skills/deckwright/runtime/deckwright-runtime.js') + '" data-hdek-runtime></script></body>')
    ctx = browser.new_context(viewport={'width': 1280, 'height': 720}); pg = ctx.new_page()
    pg.goto(url(longp)); pg.wait_for_timeout(600); pg.evaluate('scrollTo(0, 1690)'); pg.wait_for_timeout(100)
    pg.keyboard.press('z'); pg.mouse.click(620, 370); pg.wait_for_timeout(500)
    mk = pg.evaluate("(r=>[r.x+r.width/2, r.y+r.height/2])(mk.getBoundingClientRect())")
    pg.mouse.move(200, 700); pg.wait_for_timeout(200)
    barb = pg.evaluate("document.querySelector('.hdek-ui-host').shadowRoot.querySelector('.hdek-bar').getBoundingClientRect().bottom")
    pg.keyboard.press('Escape'); pg.wait_for_timeout(300)
    check('zoom on a scrolled page: the clicked point is centred, toolbar on screen, scroll restored', abs(mk[0] - 640) < 30 and abs(mk[1] - 360) < 30 and 0 < barb <= 720 and pg.evaluate('scrollY') == 1690, (mk, barb, pg.evaluate('scrollY')))
    ctx.close()

    # v0.3.5 (review): thumbnails scrolled out of view are unloaded
    manyp = os.path.join(TMP, 'many.html')
    with open(manyp, 'w') as f:
        f.write('<!doctype html><title>Many</title><main data-hdek data-width="1280" data-height="720">' + ''.join(f'<section data-slide="s{i}"><h1>Slide {i}</h1></section>' for i in range(60)) +
                '</main><script src="' + os.path.join(ROOT, 'skills/deckwright/runtime/deckwright-runtime.js') + '" data-hdek-runtime></script>')
    ctx = browser.new_context(viewport={'width': 1280, 'height': 720}); pg = ctx.new_page()
    pg.goto(url(manyp)); pg.wait_for_timeout(400); pg.keyboard.press('g'); pg.wait_for_timeout(1500)
    sr = "document.querySelector('.hdek-ui-host').shadowRoot"
    for _ in range(6): pg.evaluate(sr + ".querySelector('.hdg-cards').scrollBy(0, 400)"); pg.wait_for_timeout(500)
    nf = pg.evaluate(sr + ".querySelectorAll('.hdg-th iframe').length"); first = pg.evaluate(sr + ".querySelectorAll('.hdg-card')[0].querySelector('iframe') === null")
    check('overview keeps only thumbnails near the screen alive', nf <= 24 and first, (nf, first))
    ctx.close()

    # v0.3.5 (review): ink on a large HiDPI screen keeps full resolution
    ctx = browser.new_context(viewport={'width': 2560, 'height': 1440}, device_scale_factor=2); pg = ctx.new_page()
    pg.goto(url(os.path.join(DIST, 'demo.html'))); pg.wait_for_timeout(500)
    cw = pg.evaluate("document.querySelector('.hdek-ink').width")
    check('audience ink canvas is not capped (2560 px at 2×)', cw == 5120, cw)
    ctx.close()
    # deck CSS must not reach the runtime UI
    ctx = browser.new_context(viewport={'width': 1280, 'height': 720}); pg = ctx.new_page()
    pg.goto(url(os.path.join(DIST, 'demo.html'))); pg.wait_for_timeout(400)
    pg.add_style_tag(content='td{color:#333;font-size:28px} table{background:#fff} h3{font-size:56px} button{background:red!important}')
    pg.keyboard.press('?'); pg.wait_for_timeout(200)
    css = pg.evaluate('(()=>{const r=document.querySelector(".hdek-ui-host").shadowRoot;return [getComputedStyle(r.querySelector(".hdek-help td")).color,getComputedStyle(r.querySelector(".hdek-help h3")).fontSize]})()')
    check('deck CSS does not leak into the help card', css == ['rgb(255, 255, 255)', '17px'], css)
    pg.keyboard.press('Escape')
    # stylus writes on the first stroke without picking a tool, and the stroke does not advance the slide
    cdp = ctx.new_cdp_session(pg)
    def pen(t, x, y, b=1): cdp.send('Input.dispatchMouseEvent', {'type': t, 'x': x, 'y': y, 'button': 'left' if b else 'none', 'buttons': b, 'clickCount': 1, 'pointerType': 'pen'})
    before = pg.evaluate('DeckWright.state().slide')
    for i in range(2):
        pen('mouseMoved', 300, 300 + i * 60, 0); pen('mousePressed', 300, 300 + i * 60)
        for k in range(1, 12): pen('mouseMoved', 300 + k * 10, 300 + i * 60 + k * 3)
        pen('mouseReleased', 410, 333 + i * 60); pg.wait_for_timeout(150)
    ink = pg.evaluate('Object.values(DeckWright.ink()).reduce((a,v)=>a+v.length,0)')
    check('stylus: every stroke draws, none advances', ink == 2 and pg.evaluate('DeckWright.state().slide') == before, (ink, before))
    # PageDown works while the presenter's minutes box has focus
    with ctx.expect_page() as pi: pg.keyboard.press('s')
    pp = pi.value; pp.wait_for_timeout(1500)
    pp.click('#dur'); pp.keyboard.type('10'); a0 = pg.evaluate('JSON.stringify(DeckWright.state())')
    pp.keyboard.press('PageDown'); pp.wait_for_timeout(300)
    check('clicker (PageDown) works while a text box has focus', pg.evaluate('JSON.stringify(DeckWright.state())') != a0)
    ctx.close()
    # S with two screens (mocked): slides fullscreen on the projector + presenter popup on the laptop, one key press
    MOCK = """(() => {
      const laptop = {availLeft:0, availTop:0, availWidth:1440, availHeight:860, isInternal:true, label:'laptop'};
      const proj = {availLeft:1440, availTop:0, availWidth:1920, availHeight:1080, isInternal:false, label:'projector'};
      Object.defineProperty(window.screen, 'isExtended', {get: () => true});
      window.getScreenDetails = () => Promise.resolve({screens:[laptop, proj], currentScreen: laptop});
      const q = navigator.permissions.query.bind(navigator.permissions);
      navigator.permissions.query = (d) => d && d.name === 'window-management' ? Promise.resolve({state:'granted'}) : q(d);
      window.__calls = [];
      Element.prototype.requestFullscreen = function (o) { window.__calls.push(['fullscreen', o && o.screen && o.screen.label]); return Promise.resolve(); };
      const open = window.open.bind(window);
      window.open = function (u, n, f) { window.__calls.push(['open', f]); return open(u, n, f); };
    })();"""
    ctx = browser.new_context(viewport={'width': 1280, 'height': 720}); ctx.add_init_script(MOCK); pg = ctx.new_page()
    pg.goto(url(os.path.join(DIST, 'demo.html'))); pg.wait_for_timeout(600)
    with ctx.expect_page() as pi: pg.keyboard.press('s')
    pi.value.wait_for_timeout(500)
    calls = pg.evaluate('window.__calls')
    check('S: slides fullscreen on the projector and presenter on the laptop in one press', calls[:2] == [['fullscreen', 'projector'], ['open', 'popup=yes,left=0,top=0,width=1440,height=860']], calls)
    ctx.close()
    # a frame inside a slide cannot drive the presenter (here: clear all ink)
    ctx = browser.new_context(viewport={'width': 1280, 'height': 720}); pg = ctx.new_page()
    pg.goto(url(built('embed-post'))); pg.wait_for_timeout(500)
    pg.keyboard.press('d'); pg.mouse.move(300, 300); pg.mouse.down(); pg.mouse.move(400, 350); pg.mouse.up(); pg.keyboard.press('Escape')
    with ctx.expect_page() as pi: pg.keyboard.press('s')
    pp = pi.value; pp.wait_for_timeout(2200)
    check('postMessage from an embedded frame is ignored', len(pg.evaluate('DeckWright.ink()')) == 1)
    ctx.close()

    browser.close()

# ------------------------------------------------------------ CLI
print('CLI')
out = os.path.join(TMP, 'a.hdek')
r = cli('pack', 'test/fixtures/assets-deck', '-o', out)
check('pack succeeds', r.returncode == 0, r.stderr)
z = zipfile.ZipFile(out); names = z.namelist()
check('mimetype first and stored', names[0] == 'mimetype' and z.getinfo('mimetype').compress_type == 0, names)
check('manifest + non-ASCII (Korean) file name preserved', 'manifest.json' in names and 'img/한글 그림.png' in names, names)
check('standard unzip integrity', z.testzip() is None)
m = json.loads(z.read('manifest.json')); check('manifest format', m['format'] == 'hdek/1' and len(m['slides']) == 2, m)
r = cli('unpack', out, '-o', os.path.join(TMP, 'u'))
check('unpack succeeds', r.returncode == 0 and os.path.exists(os.path.join(TMP, 'u', 'index.html')), r.stderr)
r = cli('check', 'examples/demo'); check('check: demo has no errors', r.returncode == 0, r.stdout)
r = cli('export', 'dist/demo.hdek', '--pptx', os.path.join(TMP, 'd.pptx'), '--pdf', os.path.join(TMP, 'd.pdf'))
check('export PDF/PPTX', r.returncode == 0 and os.path.getsize(os.path.join(TMP, 'd.pptx')) > 50000, r.stderr or r.stdout)
try:
    from pptx import Presentation
    p = Presentation(os.path.join(TMP, 'd.pptx'))
    check('PPTX has 7 slides + notes', len(p.slides) == 7 and 'title slide' in p.slides[0].notes_slide.notes_text_frame.text)
except ImportError:
    pass

print('Review regressions (CLI)')
r = cli('build', 'test/fixtures/unsafe-deck/deck', '-o', os.path.join(TMP, 'unsafe.html'))
h = open(os.path.join(TMP, 'unsafe.html'), encoding='utf-8').read()
import base64
check('build never reads files outside the deck folder', 'TOPSECRET' not in h and base64.b64encode(b'TOPSECRET-DO-NOT-SHIP').decode() not in h and 'Outside the deck folder' in r.stdout, r.stdout)
# update: older DeckWright files in a folder get the current runtime, exactly as a fresh inject would
import re as _re
ud = os.path.join(TMP, 'upd'); os.makedirs(os.path.join(ud, 'sub'), exist_ok=True)
with open(os.path.join(ROOT, 'test/fixtures/media-deck/index.html'), encoding='utf-8') as f: src_html = f.read()
with open(os.path.join(ud, 'talk.html'), 'w', encoding='utf-8') as f: f.write(src_html)
cli('inject', os.path.join(ud, 'talk.html')); cli('inject', os.path.join(ud, 'talk.html'), '-o', os.path.join(ud, 'fresh.html'))
dw = os.path.join(ud, 'talk_DeckWright.html')
with open(dw, encoding='utf-8') as f: cur_html = f.read()
ver = _re.search(r'DeckWright runtime v(\d+\.\d+\.\d+)', cur_html).group(1)
with open(dw, 'w', encoding='utf-8') as f: f.write(cur_html.replace('DeckWright runtime v' + ver, 'DeckWright runtime v0.1.0'))
cli('runtime', '-o', os.path.join(ud, 'sub', 'deckwright-runtime.js'))
rj = os.path.join(ud, 'sub', 'deckwright-runtime.js')
with open(rj, encoding='utf-8') as f: rt = f.read()
with open(rj, 'w', encoding='utf-8') as f: f.write(rt.replace('runtime v' + ver, 'runtime v0.2.0'))
r1 = cli('update', ud, '--dry-run')
with open(dw, encoding='utf-8') as f: still_old = 'runtime v0.1.0' in f.read()
check('update --dry-run lists old files and writes nothing', r1.returncode == 0 and '0.1.0 → ' in r1.stdout and '0.2.0 → ' in r1.stdout and still_old, r1.stdout)
import shutil as _sh
_sh.copy(os.path.join(ROOT, 'skills/deckwright/player.html'), os.path.join(ud, 'player.html'))
with open(os.path.join(ud, 'player.html'), encoding='utf-8') as f: player_before = f.read()
r2 = cli('update', ud); r3 = cli('update', ud)
with open(os.path.join(ud, 'player.html'), encoding='utf-8') as f: player_after = f.read()
check('update never touches the player (it only mentions the runtime tag in its code)', player_before == player_after and 'player.html' not in r2.stdout, r2.stdout)
with open(dw, encoding='utf-8') as f: a_ = f.read()
with open(os.path.join(ud, 'fresh.html'), encoding='utf-8') as f: b_ = f.read()
check('update rewrites them like a fresh inject, and a second run changes nothing', r2.returncode == 0 and a_ == b_ and 'Updated 0' in r3.stdout and '3 already current' in r3.stdout, (r2.stdout, r3.stdout))
import re as _re
m = _re.search(r'id="hdek-manifest">(.*?)</script>', h, _re.S)
check('"$" in the title keeps the manifest valid', m and json.loads(m.group(1))['title'] == "Revenue ($'000)")
check('srcset keeps an existing data: URL intact', 'srcset="data:image/png;base64,iVBORw0KGgo= 1x, data:image/png;base64,' in h)
check('@import with layer() is kept as an @import', '") layer(base);' in h)
check('defer script keeps defer', '<script defer src="data:text/javascript;base64,' in h)
with sync_playwright() as pw:
    b = pw.chromium.launch(); pg = b.new_page(); errs = []
    pg.on('pageerror', lambda e: errs.append(str(e)))
    pg.goto(url(os.path.join(TMP, 'unsafe.html'))); pg.wait_for_timeout(500)
    check('built deck runs (deferred script after parsing)', not errs and pg.inner_text('h1') == 'deferred ok', (errs, pg.inner_text('h1')))
    b.close()
r = cli('pack', 'test/fixtures/unsafe-deck/deck', '-o', os.path.join(TMP, 'unsafe.hdek'))
names = zipfile.ZipFile(os.path.join(TMP, 'unsafe.hdek')).namelist()
check('pack: no ../ entries, @imported CSS assets included', not any('..' in n for n in names) and 'css/b.css' in names and 'photo (1).png' in names, names)
r = cli('build', 'test/fixtures/euckr-deck', '-o', os.path.join(TMP, 'euckr.html'))
check('EUC-KR deck converted to UTF-8', '<title>한글 발표</title>' in open(os.path.join(TMP, 'euckr.html'), encoding='utf-8').read() and 'euc-kr' in r.stdout, r.stdout)
# ZIP safety
zp = os.path.join(TMP, 'names.zip')
with zipfile.ZipFile(zp, 'w') as zf:
    zf.writestr('index.html', '<img src="이미지.png">'); zf.writestr('__proto__', 'x')
node = """
const C=require(process.argv[1]);const fs=require('fs'),z=require('zlib');
(async()=>{const out={};
 const f=await C.readZip(new Uint8Array(fs.readFileSync(process.argv[2])),(b,m)=>new Uint8Array(z.inflateRawSync(b,{maxOutputLength:Math.max(1,m)})));
 out.proto=Object.getPrototypeOf(f)===null&&Object.keys(f).includes('__proto__');
 try{await C.readZip(new Uint8Array(fs.readFileSync(process.argv[3])),(b,m)=>new Uint8Array(z.inflateRawSync(b,{maxOutputLength:Math.max(1,m)})),{entry:1e6,total:1e6,files:100});out.bomb='read'}catch(e){out.bomb=e.message}
 try{C.writeZip(Array.from({length:65537},(_, i)=>({name:'f'+i,bytes:new Uint8Array(1)})));out.many='ok'}catch(e){out.many=e.message}
 console.log(JSON.stringify(out));})();
"""
bomb = os.path.join(TMP, 'bomb.zip')
with zipfile.ZipFile(bomb, 'w', zipfile.ZIP_DEFLATED) as zb: zb.writestr('index.html', 'x'); zb.writestr('big.bin', b'\0' * (8 * 1024 * 1024))
o = subprocess.run(['node', '-e', node, os.path.join(ROOT, 'skills/deckwright/core/deckwright-core.js'), zp, bomb], capture_output=True, text=True)
res = json.loads(o.stdout or '{}')
check('ZIP: "__proto__" entry cannot change the file map', res.get('proto') is True, o.stdout + o.stderr)
check('ZIP: size limits stop a ZIP bomb', 'expands' in str(res.get('bomb')), res)
check('ZIP: more than 65,535 files is refused, not corrupted', 'Too many' in str(res.get('many')), res)
# an entry name written in CP949 without the UTF-8 flag (Korean Windows zip tools)
cpz = os.path.join(TMP, 'cp949name.zip')
with zipfile.ZipFile(cpz, 'w') as zc:
    zc.writestr('index.html', '<img src="이미지.png">'); zc.writestr('abcdef.png', b'PNG')
raw = open(cpz, 'rb').read().replace(b'abcdef.png', '이미지.png'.encode('cp949')); open(cpz, 'wb').write(raw)
r = cli('build', cpz, '-o', os.path.join(TMP, 'cp949.html'))
check('ZIP: CP949 file names are read correctly', r.returncode == 0 and 'Missing' not in r.stdout and 'data:image/png;base64,UE5H' in open(os.path.join(TMP, 'cp949.html'), encoding='utf-8').read(), r.stdout + r.stderr)
r = cli('biuld', 'x'); check('unknown command exits 1', r.returncode == 1)
import time as _t
t0 = _t.time(); r = subprocess.run(['node', os.path.join(ROOT, 'skills', 'deckwright', 'scripts', 'deckwright.mjs'), 'export', 'dist/demo.html', '--pdf'], cwd=ROOT, capture_output=True, text=True, timeout=60)
check('export with a missing value fails fast (no hang)', r.returncode == 1 and _t.time() - t0 < 20, r.stderr)
# inject (the default path in SKILL.md)
src = os.path.join(TMP, 'inj'); os.makedirs(src, exist_ok=True)
open(os.path.join(src, 'talk.html'), 'w').write(open(os.path.join(ROOT, 'examples/ai-style/index.html')).read())
orig = open(os.path.join(src, 'talk.html')).read()
r = cli('inject', os.path.join(src, 'talk.html'))
outp = os.path.join(src, 'talk_DeckWright.html')
check('inject writes <name>_DeckWright.html and leaves the original alone', os.path.exists(outp) and 'data-hdek-runtime' in open(outp).read() and open(os.path.join(src, 'talk.html')).read() == orig, r.stdout + r.stderr)
check('unpack keeps the Korean file name', os.path.exists(os.path.join(TMP, 'u', 'img', '한글 그림.png')))
check('export wrote a 7-page PDF', os.path.exists(os.path.join(TMP, 'd.pdf')) and __import__('pdfplumber').open(os.path.join(TMP, 'd.pdf')).pages.__len__() == 7)

print('Web fonts (--fetch-fonts)')
import http.server, socketserver, threading, functools
wf = os.path.join(TMP, 'webfont-host'); os.makedirs(wf, exist_ok=True)
open(os.path.join(wf, 'a.woff2'), 'wb').write(b'wOF2-hangul-slice')
open(os.path.join(wf, 'b.woff2'), 'wb').write(b'wOF2-cyrillic-slice')
open(os.path.join(wf, 'css2'), 'w').write("""@font-face { font-family: 'T'; font-weight: 400; src: url(a.woff2) format('woff2'); unicode-range: U+AC00-D7A3; }
@font-face { font-family: 'T'; font-weight: 700; src: url(a.woff2) format('woff2'); unicode-range: U+AC00-D7A3; }
@font-face { font-family: 'T'; font-weight: 400; src: url(b.woff2) format('woff2'); unicode-range: U+0400-045F; }
""")
class Quiet(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a): pass
srv = socketserver.TCPServer(('127.0.0.1', 0), functools.partial(Quiet, directory=wf)); port = srv.server_address[1]
threading.Thread(target=srv.serve_forever, daemon=True).start()
fd = os.path.join(TMP, 'webfont-deck'); os.makedirs(os.path.join(fd, 'css'), exist_ok=True)
css_url = f'http://127.0.0.1:{port}/css2?family=T:wght@400;700&amp;display=swap'
open(os.path.join(fd, 'index.html'), 'w').write(f'<!doctype html><html><head><meta charset="utf-8"><link rel="preconnect" href="https://fonts.googleapis.com"><link rel="stylesheet" href="{css_url}"><link rel="stylesheet" href="css/x.css"><style>body{{font-family:T}}</style></head><body><main data-hdek><section data-slide="a"><h1>가나다</h1></section></main></body></html>')
open(os.path.join(fd, 'css', 'x.css'), 'w').write(f'@import url("http://127.0.0.1:{port}/css2?family=T:wght@400");\nh1{{color:red}}')
r = cli('build', fd, '--fetch-fonts', '-o', os.path.join(TMP, 'wf.html'))
h = open(os.path.join(TMP, 'wf.html'), encoding='utf-8').read()
check('build --fetch-fonts embeds the web font', 'data:font/woff2;base64,' + base64.b64encode(b'wOF2-hangul-slice').decode() in h and 'Web fonts embedded' in r.stdout, r.stdout + r.stderr)
check('only slices for characters in the deck are kept', base64.b64encode(b'wOF2-cyrillic-slice').decode() not in h)
check('one font file per slice (weights merged into a range)', h.count(base64.b64encode(b'wOF2-hangul-slice').decode()) == 2 and 'font-weight: 400 700' in h, h.count(base64.b64encode(b'wOF2-hangul-slice').decode()))
check('no request to the font host is left', not _re.search(r'(href|url)\(?=?["\']?http://127\.0\.0\.1', h) and 'rel="preconnect"' not in h)
r = cli('build', fd, '-o', os.path.join(TMP, 'wf0.html'))
check('without --fetch-fonts: nothing is downloaded and a hint is shown', '--fetch-fonts' in r.stdout and 'data:font/woff2' not in open(os.path.join(TMP, 'wf0.html'), encoding='utf-8').read(), r.stdout)
r = cli('inject', os.path.join(fd, 'index.html'), '--fetch-fonts')
inj = open(os.path.join(fd, 'index_DeckWright.html'), encoding='utf-8').read()
check('inject --fetch-fonts embeds fonts and keeps local references', 'data:font/woff2' in inj and 'href="css/x.css"' in inj, r.stdout + r.stderr)
r = cli('pack', fd, '--fetch-fonts', '-o', os.path.join(TMP, 'wf.hdek'))
zx = zipfile.ZipFile(os.path.join(TMP, 'wf.hdek'))
check('pack --fetch-fonts rewrites the page and its CSS files', 'data:font/woff2' in zx.read('index.html').decode() and 'data:font/woff2' in zx.read('css/x.css').decode(), r.stdout + r.stderr)
srv.shutdown()

ok = sum(1 for r in results if r[1]); total = len(results)
print(f'\n{ok}/{total} passed')
sys.exit(0 if ok == total else 1)
